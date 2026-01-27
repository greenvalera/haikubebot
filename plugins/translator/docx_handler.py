"""
DOCX file translation handler.
Adapted from mstranslate package.
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openai import OpenAI

from .translator import translate_texts


W_R = qn("w:r")
W_T = qn("w:t")
W_P = qn("w:p")
W_TXBX_CONTENT = qn("w:txbxContent")
W_SDT_CONTENT = qn("w:sdtContent")


def _get_all_runs_in_paragraph(paragraph):
    """Get ALL w:r elements in a paragraph, including those inside
    w:hyperlink, w:sdt, w:smartTag, etc."""
    results = []
    p_elem = paragraph._p

    for r_elem in p_elem.iter(W_R):
        parent = r_elem.getparent()
        while parent is not None and parent.tag != W_P:
            parent = parent.getparent()
        if parent is not p_elem:
            continue

        t_elems = r_elem.findall(W_T)
        text = "".join(t.text or "" for t in t_elems)
        if text and text.strip():
            results.append((text, r_elem, t_elems))
    return results


def _set_run_text(t_elems, new_text):
    """Write new text back into w:t elements."""
    if len(t_elems) == 1:
        t_elems[0].text = new_text
        if new_text and (new_text[0] == " " or new_text[-1] == " "):
            t_elems[0].set(qn("xml:space"), "preserve")
    elif t_elems:
        t_elems[0].text = new_text
        if new_text and (new_text[0] == " " or new_text[-1] == " "):
            t_elems[0].set(qn("xml:space"), "preserve")
        for t in t_elems[1:]:
            t.text = ""


def _collect_from_paragraphs(paragraphs, segments, seen_run_ids):
    """Collect text runs from a list of python-docx Paragraph objects."""
    for para in paragraphs:
        for text, r_elem, t_elems in _get_all_runs_in_paragraph(para):
            r_id = id(r_elem)
            if r_id in seen_run_ids:
                continue
            seen_run_ids.add(r_id)
            segments.append((text, r_elem, t_elems))


def _collect_textbox_paragraphs(root_element):
    """Find paragraphs inside text boxes (w:txbxContent) anywhere in the XML tree."""
    from docx.text.paragraph import Paragraph

    paragraphs = []
    for txbx in root_element.iter(W_TXBX_CONTENT):
        for p_elem in txbx.iter(W_P):
            paragraphs.append(Paragraph(p_elem, None))
    return paragraphs


def translate_docx(
    input_path: Path,
    output_path: Path,
    source_lang: str,
    target_lang: str,
    client: OpenAI,
    model: str = "gpt-5.2",
) -> None:
    """Translate all text in a .docx file."""
    doc = Document(str(input_path))

    segments = []
    seen_run_ids = set()

    # 1. Body paragraphs
    _collect_from_paragraphs(doc.paragraphs, segments, seen_run_ids)

    # 2. SDT blocks (e.g. Table of Contents)
    from docx.text.paragraph import Paragraph

    for sdt_content in doc.element.body.iter(W_SDT_CONTENT):
        for p_elem in sdt_content.iter(W_P):
            sdt_para = Paragraph(p_elem, None)
            _collect_from_paragraphs([sdt_para], segments, seen_run_ids)

    # 3. Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _collect_from_paragraphs(cell.paragraphs, segments, seen_run_ids)

    # 4. Headers and footers
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if hf is not None:
                _collect_from_paragraphs(hf.paragraphs, segments, seen_run_ids)
                for table in hf.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            _collect_from_paragraphs(
                                cell.paragraphs, segments, seen_run_ids
                            )

    # 5. Text boxes (captions, callouts, etc.)
    textbox_paras = _collect_textbox_paragraphs(doc.element)
    _collect_from_paragraphs(textbox_paras, segments, seen_run_ids)
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if hf is not None:
                tb_paras = _collect_textbox_paragraphs(hf._element)
                _collect_from_paragraphs(tb_paras, segments, seen_run_ids)

    if not segments:
        doc.save(str(output_path))
        print("  No text found, saved as-is.")
        return

    texts = [text for text, _, _ in segments]
    print(f"  Found {len(texts)} text segments to translate...")

    translated = translate_texts(texts, source_lang, target_lang, client, model)

    for (_, r_elem, t_elems), new_text in zip(segments, translated):
        _set_run_text(t_elems, new_text)

    doc.save(str(output_path))
