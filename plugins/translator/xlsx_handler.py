"""
XLSX file translation handler.
Adapted from mstranslate package.
"""

from pathlib import Path

from openpyxl import load_workbook
from openai import OpenAI

from .translator import translate_texts


def translate_xlsx(
    input_path: Path,
    output_path: Path,
    source_lang: str,
    target_lang: str,
    client: OpenAI,
    model: str = "gpt-5.2",
) -> None:
    """Translate all text in an .xlsx file."""
    wb = load_workbook(str(input_path))

    cell_refs: list[tuple] = []
    sheet_names: list[str] = []

    for sheet in wb.worksheets:
        sheet_names.append(sheet.title)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    cell_refs.append((sheet, cell.row, cell.column, cell.value))

    cell_texts = [val for _, _, _, val in cell_refs]
    all_texts = cell_texts + sheet_names

    if not all_texts or not any(t.strip() for t in all_texts):
        wb.save(str(output_path))
        print("  No text found, saved as-is.")
        return

    print(
        f"  Found {len(cell_texts)} text cells and {len(sheet_names)} sheets to translate..."
    )

    translated = translate_texts(all_texts, source_lang, target_lang, client, model)

    for i, (sheet, row, col, _) in enumerate(cell_refs):
        sheet.cell(row=row, column=col).value = translated[i]

    for i, sheet in enumerate(wb.worksheets):
        new_name = translated[len(cell_texts) + i]
        sheet.title = new_name

    wb.save(str(output_path))
